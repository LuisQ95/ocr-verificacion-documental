"""Salida del pipeline: base SQLite (para consultas) y reporte Excel (para revisión humana)."""

import json
import sqlite3
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import config as C
from .comparacion import (AUTENTICO, DESACTUALIZADO, DESACTUALIZADO_C, DIFERENTE, DUDOSO, FALSIFICADO, IGUAL,
                          NO_APLICA, OBSERVADO)
from .chile.requisitos_cl import APTO, INCOMPLETO, NO_APTO, NO_EVALUABLE

COLORES = {FALSIFICADO: "F4C7C3", OBSERVADO: "FCE8B2", AUTENTICO: "B7E1CD", DESACTUALIZADO: "C9DAF8",
           DIFERENTE: "F4C7C3", DUDOSO: "FCE8B2", DESACTUALIZADO_C: "C9DAF8",
           APTO: "B7E1CD", NO_APTO: "F4C7C3", INCOMPLETO: "FCE8B2", NO_EVALUABLE: "FCE8B2",
           "NO LEGIBLE": "FCE8B2", "NO VÁLIDO": "C9DAF8", "NO RECONOCIDO": "EEEEEE", "SIN DATO OFICIAL": "EEEEEE"}


def _sqlite(ruta, personas, documentos, meta, no_admitidos):
    """Base SQLite (se rehace desde cero) con la ejecución, personas, documentos, campos, señales y no admitidos."""
    ruta.unlink(missing_ok=True)
    con = sqlite3.connect(ruta)
    con.executescript("""
        CREATE TABLE ejecucion (clave TEXT PRIMARY KEY, valor TEXT);
        CREATE TABLE personas (carpeta TEXT PRIMARY KEY, dni TEXT, declarado INTEGER, veredicto TEXT,
            documentos INTEGER, falsificados INTEGER, observados INTEGER, desactualizados INTEGER,
            tipos_falsificados TEXT, puesto TEXT, aptitud TEXT, motivos_aptitud TEXT);
        CREATE TABLE documentos (carpeta TEXT, archivo TEXT, tipo_documento TEXT, formato TEXT, veredicto TEXT,
            motivos TEXT, misma_instancia INTEGER, archivo_oficial TEXT, alineacion_ok INTEGER, alineacion TEXT,
            residual_etiquetas_pt REAL, producer TEXT, PRIMARY KEY (carpeta, archivo));
        CREATE TABLE campos (carpeta TEXT, archivo TEXT, zona TEXT, oficial TEXT, postulante TEXT, estado TEXT,
            motivo TEXT, conf REAL);
        CREATE TABLE senales (carpeta TEXT, archivo TEXT, tipo TEXT, severidad TEXT, detalle TEXT);
        CREATE TABLE no_admitidos (carpeta TEXT, archivo TEXT, origen TEXT, motivo TEXT);
    """)
    con.executemany("INSERT INTO ejecucion VALUES (?, ?)", [(k, json.dumps(v, ensure_ascii=False)) for k, v in meta.items()])
    con.executemany("INSERT INTO personas VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    [(p["carpeta"], p["dni"], int(p["declarado"]), p["veredicto"], p["documentos"], p["falsificados"],
                      p["observados"], p["desactualizados"], ", ".join(p["tipos_falsificados"]), p["puesto"],
                      p["aptitud"], " | ".join(p["motivos_aptitud"])) for p in personas])
    for d in documentos:
        con.execute("INSERT INTO documentos VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    (d["carpeta"], d["archivo"], d["tipo_documento"], d["formato"], d["veredicto"],
                     " | ".join(d["motivos"]), None if d["misma_instancia"] is None else int(d["misma_instancia"]),
                     d["archivo_oficial"], int(bool(d["alineacion_ok"])),
                     json.dumps(d["alineacion"], ensure_ascii=False), d["residual_etiquetas_pt"], d["producer"]))
        con.executemany("INSERT INTO campos VALUES (?,?,?,?,?,?,?,?)",
                        [(d["carpeta"], d["archivo"], c["zona"], c["oficial"], c["postulante"], c["estado"],
                          c["motivo"], c["conf"]) for c in d["campos"]])
        con.executemany("INSERT INTO senales VALUES (?,?,?,?,?)",
                        [(d["carpeta"], d["archivo"], s["tipo"], s["severidad"], s["detalle"]) for s in d["senales"]])
    con.executemany("INSERT INTO no_admitidos VALUES (?,?,?,?)",
                    [(x["carpeta"], x["archivo"], x["origen"], x["motivo"]) for x in no_admitidos])
    con.commit()
    con.close()


def _hoja(ws, encabezados, filas, cols_color=()):
    """Llena una hoja: encabezado con estilo, filas, color por estado en `cols_color`, filtro y anchos."""
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1A264D")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in filas:
        ws.append(fila)
        for i in cols_color:
            color = COLORES.get(fila[i])
            if color:
                ws.cell(ws.max_row, i + 1).fill = PatternFill("solid", fgColor=color)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ancho = max(len(str(c.value or "")) for c in col[:300])
        ws.column_dimensions[col[0].column_letter].width = min(70, max(10, ancho + 2))


def _conteo(ws, titulo, contador):
    """Agrega a la hoja una tabla de conteo (valor, cantidad) coloreada por valor."""
    ws.append([])
    ws.append([titulo, "Cantidad"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for v, n in contador.most_common():
        ws.append([v, n])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=COLORES.get(v, "FFFFFF"))


def _excel(ruta, personas, documentos, meta, no_admitidos):
    """Reporte Excel: resumen con guía de lectura y conteos, y hojas de personas, documentos, diferencias, señales
    y no admitidos."""
    wb = Workbook()
    ws = wb.active
    ws.title = "resumen"
    ws.append(["Reporte de verificación de documentos"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for clave, texto in (
            ("Cómo leer este reporte", ""),
            ("Autenticidad", "¿Los documentos son verdaderos? AUTÉNTICO = coinciden con el dato oficial (JSON); "
                             "FALSIFICADO = algún dato o elemento no coincide; OBSERVADO = requiere revisión humana "
                             "(lectura dudosa); NO LEGIBLE = no se pudo leer (no se acusa); NO VÁLIDO = copia no "
                             "válida como certificado; DESACTUALIZADO = genuino pero antiguo; NO RECONOCIDO = archivo "
                             "que no es ninguno de los documentos solicitados; SIN DATO OFICIAL = no se entregó su "
                             "JSON oficial: no se verificó, solo se extrajeron sus datos."),
            ("Aptitud", "¿Cumple los requisitos del puesto (config/requisitos_cl.yaml)? Se evalúa con los datos "
                        "oficiales. APTO / NO APTO / INCOMPLETO (falta un documento o el presentado es falso). "
                        "Un postulante puede ser AUTÉNTICO y NO APTO: sus documentos son verdaderos pero no "
                        "cumple un requisito (p. ej. licencia vencida)."),
            ("Detalle", "Hoja 'personas': motivos de aptitud. Hoja 'diferencias': campo por campo, oficial vs "
                        "postulante. Hoja 'senales': evidencia adicional (QR, objetos superpuestos, identidad...).")):
        ws.append([clave, texto])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])
    for clave, valor in meta.items():
        ws.append([clave, json.dumps(valor, ensure_ascii=False) if isinstance(valor, (dict, list)) else valor])
    _conteo(ws, "Autenticidad por postulante", Counter(p["veredicto"] for p in personas))
    _conteo(ws, "Aptitud por postulante", Counter(p["aptitud"] for p in personas))
    _conteo(ws, "Veredicto por documento", Counter(d["veredicto"] for d in documentos))
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90

    _hoja(wb.create_sheet("personas"),
          ["Carpeta", "RUT / documento de identidad", "Identidad declarada", "Autenticidad", "Aptitud", "Motivos de aptitud",
           "Documentos",
           "Falsificados", "Observados", "Desactualizados", "Documentos falsificados", "Puesto"],
          [[p["carpeta"], p["dni"], "SÍ" if p["declarado"] else "NO", p["veredicto"], p["aptitud"],
            " | ".join(p["motivos_aptitud"]), p["documentos"], p["falsificados"], p["observados"],
            p["desactualizados"], ", ".join(p["tipos_falsificados"]), p["puesto"]] for p in personas],
          cols_color=(3, 4))

    _hoja(wb.create_sheet("documentos"),
          ["Carpeta", "Archivo", "Tipo de documento", "Formato", "Veredicto", "Motivos", "Mismo documento que el oficial",
           "Archivo oficial"],
          [[d["carpeta"], d["archivo"], d["tipo_documento"], d["formato"], d["veredicto"],
            " | ".join(d["motivos"])[:2000],
            "" if d["misma_instancia"] is None else ("SÍ" if d["misma_instancia"] else "NO (otra emisión)"),
            d["archivo_oficial"]] for d in documentos], cols_color=(4,))

    _hoja(wb.create_sheet("diferencias"),
          ["Carpeta", "Archivo", "Campo", "Valor oficial", "Valor postulante", "Estado", "Motivo", "Confianza OCR"],
          [[d["carpeta"], d["archivo"], c["zona"], c["oficial"], c["postulante"], c["estado"], c["motivo"], c["conf"]]
           for d in documentos for c in d["campos"] if c["estado"] not in (IGUAL, NO_APLICA)], cols_color=(5,))

    _hoja(wb.create_sheet("senales"),
          ["Carpeta", "Archivo", "Tipo", "Severidad", "Detalle"],
          [[d["carpeta"], d["archivo"], s["tipo"], s["severidad"], s["detalle"]]
           for d in documentos for s in d["senales"]])

    _hoja(wb.create_sheet("no_admitidos"), ["Carpeta", "Archivo", "Origen", "Motivo"],
          [[x["carpeta"], x["archivo"], x["origen"], x["motivo"]] for x in no_admitidos])
    wb.save(ruta)


def escribir_reporte(personas, documentos, meta, salida=C.DIR_RESULTADOS, no_admitidos=()):
    """Escribe verificacion.db y reporte_verificacion.xlsx en `salida`; devuelve sus rutas."""
    salida.mkdir(parents=True, exist_ok=True)
    db = salida / "verificacion.db"
    xlsx = salida / "reporte_verificacion.xlsx"
    _sqlite(db, personas, documentos, meta, list(no_admitidos))
    _excel(xlsx, personas, documentos, meta, list(no_admitidos))
    return db, xlsx
